import collections
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_variant_1():
    prs = Presentation()
    
    # Slide 1: Title
    slide_layout = prs.slide_layouts[0] # Title slide
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Smart Attendance Tracking"
    subtitle.text = "PSE Surat College\nGTU Mandatory 75% Attendance Criteria"
    
    try:
        img_path = r"c:\Users\Lucifer\OneDrive\Desktop\attendance\pse_real.png"
        slide.shapes.add_picture(img_path, Inches(3), Inches(4.5), height=Inches(2.5))
    except Exception as e:
        print("Image 1 not found:", e)

    # Slide 2: Problem
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "The Problem: Old Pen & Paper Format"
    
    content = slide.shapes.placeholders[1].text_frame
    content.text = "Earlier, attendance was taken manually, which caused several issues:"
    p = content.add_paragraph()
    p.text = "As a Student: Unsure of current attendance %, proxies by others, difficult to track GTU 75% criteria."
    p = content.add_paragraph()
    p.text = "As a Teacher: Time-consuming to call out names, manual data entry in excel, hard to calculate low attendance."
    
    try:
        img_path = r"C:\Users\Lucifer\.gemini\antigravity\brain\9594b305-33ed-430c-8df0-2b2ff4b4b3de\excel_attendance_sheet_1782458816041.png"
        slide.shapes.add_picture(img_path, Inches(6.5), Inches(0.5), width=Inches(3))
    except Exception as e:
        print("Image 2 not found:", e)

    # Slide 3: Solution
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Our Solution"
    
    content = slide.shapes.placeholders[1].text_frame
    content.text = "Smart Attendance Tracking site digitizes the entire process."
    p = content.add_paragraph()
    p.text = "Student Solution: Real-time dashboard to check attendance %, alerts for low attendance, no proxies."
    p = content.add_paragraph()
    p.text = "Teacher Solution: One-click attendance taking, automated calculations, export reports easily."
    
    try:
        img_path = r"C:\Users\Lucifer\.gemini\antigravity\brain\9594b305-33ed-430c-8df0-2b2ff4b4b3de\smart_attendance_site_1782458828819.png"
        slide.shapes.add_picture(img_path, Inches(5), Inches(3.5), width=Inches(4.5))
    except Exception as e:
        print("Image 3 not found:", e)

    # Slide 4: Features
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Features"
    
    content = slide.shapes.placeholders[1].text_frame
    content.text = "Student Features:"
    p = content.add_paragraph()
    p.text = "- Live Attendance Dashboard\n- GTU 75% Alert System\n- Leave Request Portal"
    
    p = content.add_paragraph()
    p.text = "Teacher Features:"
    p = content.add_paragraph()
    p.text = "- Quick Mark Interface\n- Defaulter List Generator\n- Subject-wise Reports"
    
    # Slide 5: Tech Used
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Tech Used"
    
    content = slide.shapes.placeholders[1].text_frame
    content.text = "We built this with modern and easy tools:"
    p = content.add_paragraph()
    p.text = "Frontend: HTML, CSS, JavaScript (The face of the site, e.g., buttons and colors)"
    p = content.add_paragraph()
    p.text = "Backend: Node.js / Python (The brain behind the site doing calculations)"
    p = content.add_paragraph()
    p.text = "Database: MongoDB / SQL (Where all student data is saved securely)"
    
    # Slide 6: Any Questions
    slide_layout = prs.slide_layouts[0] # Title slide for centered text
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Any Questions?"
    # Remove subtitle
    sp = slide.placeholders[1].element
    sp.getparent().remove(sp)

    # Slide 7: URL
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Visit Our Site"
    
    content = slide.shapes.placeholders[1].text_frame
    content.text = "Click below to track your attendance now!"
    p = content.add_paragraph()
    p.text = "https://smartattendance.psesurat.edu"
    
    try:
        img_path = r"C:\Users\Lucifer\.gemini\antigravity\brain\9594b305-33ed-430c-8df0-2b2ff4b4b3de\smart_attendance_site_1782458828819.png"
        slide.shapes.add_picture(img_path, Inches(2), Inches(3), width=Inches(6))
    except Exception as e:
        print("Image 3 not found:", e)

    prs.save("Smart_Attendance_Variant1.pptx")
    print("Variant 1 generated successfully.")

def create_variant_2():
    prs = Presentation()
    
    # Apply a nice background color or theme via python-pptx is limited without a template,
    # but we can format text nicely.
    
    # Slide 1: Title
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    
    # Set a background color for the slide (Dark Blue)
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(10, 25, 47)
    
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "SMART ATTENDANCE TRACKING"
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(100, 255, 218)
    title.text_frame.paragraphs[0].font.bold = True
    
    subtitle.text = "Empowering PSE Surat College | Meeting the GTU 75% Mandate"
    subtitle.text_frame.paragraphs[0].font.color.rgb = RGBColor(204, 214, 246)
    
    try:
        img_path = r"c:\Users\Lucifer\OneDrive\Desktop\attendance\pse_real.png"
        slide.shapes.add_picture(img_path, Inches(3.5), Inches(4.5), height=Inches(2.5))
    except Exception as e:
        pass

    # Slide 2: The Problem
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(240, 242, 245)
    
    title = slide.shapes.title
    title.text = "The Challenge: Archaic Pen & Paper"
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(17, 34, 64)
    title.text_frame.paragraphs[0].font.bold = True
    
    content = slide.shapes.placeholders[1].text_frame
    content.text = "Traditional attendance tracking is prone to errors and inefficiency:"
    p = content.add_paragraph()
    p.text = "For Students: \n- Constant anxiety about meeting the 75% GTU criteria.\n- Vulnerable to fake proxies."
    p = content.add_paragraph()
    p.text = "For Teachers: \n- Wastes 10+ minutes of lecture time.\n- Cumbersome data entry at the end of the month."
    
    try:
        img_path = r"C:\Users\Lucifer\.gemini\antigravity\brain\9594b305-33ed-430c-8df0-2b2ff4b4b3de\excel_attendance_sheet_1782458816041.png"
        slide.shapes.add_picture(img_path, Inches(6.5), Inches(0.5), width=Inches(3))
    except Exception as e:
        pass

    # Slide 3: The Solution
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    
    title = slide.shapes.title
    title.text = "The Innovation: Smart Dashboard"
    title.text_frame.paragraphs[0].font.bold = True
    
    content = slide.shapes.placeholders[1].text_frame
    content.text = "A seamless digital experience replacing spreadsheets."
    p = content.add_paragraph()
    p.text = "Student Portal: \n- Crystal-clear analytics of attendance.\n- Instant alerts before falling below 75%."
    p = content.add_paragraph()
    p.text = "Teacher Portal: \n- Rapid, secure attendance marking.\n- One-click PDF/Excel report generation."
    
    try:
        img_path = r"C:\Users\Lucifer\.gemini\antigravity\brain\9594b305-33ed-430c-8df0-2b2ff4b4b3de\smart_attendance_site_1782458828819.png"
        slide.shapes.add_picture(img_path, Inches(5.5), Inches(3), width=Inches(4))
    except Exception as e:
        pass

    # Slide 4: Features
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Key Features"
    title.text_frame.paragraphs[0].font.bold = True
    
    content = slide.shapes.placeholders[1].text_frame
    content.text = "Empowering Students:"
    p = content.add_paragraph()
    p.text = "- Dynamic Attendance Gauge\n- GTU Threshold Notifications\n- Digital Leave Management"
    
    p = content.add_paragraph()
    p.text = "\nEmpowering Faculty:"
    p = content.add_paragraph()
    p.text = "- Biometric / QR Based Quick Mark\n- Auto-generated Defaulter Lists\n- Comprehensive Analytics"

    # Slide 5: Tech Used
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Technology Stack"
    title.text_frame.paragraphs[0].font.bold = True
    
    content = slide.shapes.placeholders[1].text_frame
    content.text = "Built for Scale and Speed:"
    p = content.add_paragraph()
    p.text = "User Interface (React/Next.js): Delivering a smooth, app-like experience."
    p = content.add_paragraph()
    p.text = "Server Engine (Node.js/Express): Fast processing of thousands of records."
    p = content.add_paragraph()
    p.text = "Database (MongoDB): Secure, scalable storage for college data."

    # Slide 6: Any Questions
    slide_layout = prs.slide_layouts[0] 
    slide = prs.slides.add_slide(slide_layout)
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(10, 25, 47)
    
    title = slide.shapes.title
    title.text = "Questions?"
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(100, 255, 218)
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.size = Pt(64)
    
    sp = slide.placeholders[1].element
    sp.getparent().remove(sp)

    # Slide 7: URL
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Experience It Now"
    title.text_frame.paragraphs[0].font.bold = True
    
    content = slide.shapes.placeholders[1].text_frame
    content.text = "Join the future of education."
    p = content.add_paragraph()
    p.text = "https://smartattendance.psesurat.edu"
    p.font.color.rgb = RGBColor(0, 0, 255)
    
    try:
        img_path = r"C:\Users\Lucifer\.gemini\antigravity\brain\9594b305-33ed-430c-8df0-2b2ff4b4b3de\smart_attendance_site_1782458828819.png"
        slide.shapes.add_picture(img_path, Inches(2), Inches(3), width=Inches(6))
    except Exception as e:
        pass

    prs.save("Smart_Attendance_Variant2.pptx")
    print("Variant 2 generated successfully.")

if __name__ == '__main__':
    create_variant_1()
    create_variant_2()
